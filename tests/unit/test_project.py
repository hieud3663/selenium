"""Offline regressions for false passes, isolation, data parity and reporting."""
import copy
import hashlib
import io
import json
import tempfile
import unittest
import shutil
from dataclasses import replace
from pathlib import Path
from unittest.mock import Mock, patch

from selenium.common.exceptions import TimeoutException, WebDriverException
from openpyxl import load_workbook

from testcases.reader import FIELDS, WORKBOOK_PATH, coverage_rows, load_cases, validate_cases, validate_mapping
from config import Settings
from pages.login_page import LoginPage
from reporting.test_result import EvidenceResult, result_exit_code, save_results
from tests import test_login as browser_tests


class ProjectRegressionTests(unittest.TestCase):
    def browser_case(self, id):
        case = browser_tests.TestLogin("test_" + id)
        case.settings = Settings()
        case.spec = browser_tests.CASES[id]
        case.login_page = Mock()
        case.driver = Mock()
        return case

    def test_excel_and_mapping_have_all_original_ids(self):
        cases = load_cases()
        validate_mapping(cases)
        self.assertEqual(len(cases), 40)
        self.assertTrue(all(c["status"] == "Not Run" and not c["actual"] for c in cases))

    def test_duplicate_id_or_hardcoded_pass_is_rejected(self):
        for mutation in ("duplicate", "pass"):
            with self.subTest(mutation=mutation):
                cases = copy.deepcopy(load_cases())
                if mutation == "duplicate":
                    cases[1]["id"] = cases[0]["id"]
                else:
                    cases[0]["status"] = "PASS"
                with self.assertRaises(ValueError):
                    validate_cases(cases)

    def test_excel_is_read_without_modification_and_contains_all_cases(self):
        before = hashlib.sha256(WORKBOOK_PATH.read_bytes()).hexdigest()
        cases = load_cases()
        self.assertEqual(len(cases), 40)
        self.assertEqual(list(cases[0]), list(FIELDS))
        self.assertIn("\n", cases[0]["steps"])
        self.assertTrue(all(isinstance(case["parameters"], dict) for case in cases))
        self.assertEqual(hashlib.sha256(WORKBOOK_PATH.read_bytes()).hexdigest(), before)

    def test_reviewed_excel_edits_are_read_directly_without_json_import(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "cases.xlsx"
            shutil.copyfile(WORKBOOK_PATH, path)
            workbook = load_workbook(path)
            workbook["Testcases"]["D2"] = "Reviewed title from Excel"
            rows = list(workbook["Testcases"].iter_rows())
            payload_row = next(row for row in rows if row[0].value == "TC_SEC_02")
            payload_row[-1].value = json.dumps({"payload": "EXCEL_PAYLOAD"})
            workbook.save(path)
            workbook.close()
            cases = load_cases(path)
            self.assertEqual(cases[0]["title"], "Reviewed title from Excel")
            case = self.browser_case("TC_SEC_02")
            case.spec = next(spec for spec in cases if spec["id"] == "TC_SEC_02")
            case.assert_rejected = Mock()
            case.login_page.get_page_source.return_value = "clean response"
            case.test_TC_SEC_02()
            calls = case.login_page.login.call_args_list
            self.assertEqual(calls[0].args[0], "EXCEL_PAYLOAD")
            self.assertEqual(len(calls), 1)
            self.assertFalse((Path(directory) / "testcases.json").exists())

    def test_invalid_excel_header_parameters_formula_or_status_are_rejected(self):
        for mutation in ("header", "json", "formula", "pass", "duplicate"):
            with self.subTest(mutation=mutation), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "cases.xlsx"
                shutil.copyfile(WORKBOOK_PATH, path)
                workbook = load_workbook(path)
                sheet = workbook["Testcases"]
                if mutation == "header":
                    sheet["A1"] = "Wrong header"
                elif mutation == "json":
                    sheet["R2"] = "not JSON"
                elif mutation == "formula":
                    sheet["D2"] = "=1+1"
                elif mutation == "pass":
                    sheet["K2"] = "PASS"
                else:
                    sheet["A3"] = sheet["A2"].value
                workbook.save(path)
                workbook.close()
                with self.assertRaises(ValueError):
                    load_cases(path)

    def test_execution_workbook_is_separate_and_preserves_unexecuted_cases(self):
        before = hashlib.sha256(WORKBOOK_PATH.read_bytes()).hexdigest()
        cases = load_cases()
        report = {"run_id": "offline-fixture", "results": {"TC_FUNC_01": {"status": "SKIP", "detail": "No browser"}}}
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            save_results(output, cases, report)
            workbook = load_workbook(output / "results.xlsx")
            self.assertEqual(workbook.sheetnames, ["Execution", "Coverage"])
            self.assertEqual(workbook["Execution"].max_row, 41)
            self.assertEqual(workbook["Execution"]["B2"].value, "SKIP")
            self.assertEqual(workbook["Execution"]["B3"].value, "Not Run")
            rows = [[value if value is not None else "" for value in row]
                    for row in workbook["Coverage"].iter_rows(min_row=2, values_only=True)]
            self.assertEqual(rows, coverage_rows(cases))
            workbook.close()
            self.assertEqual(json.loads((output / "results.json").read_text(encoding="utf-8")), report)
        self.assertEqual(hashlib.sha256(WORKBOOK_PATH.read_bytes()).hexdigest(), before)

    def test_wrong_credentials_cannot_pass_without_expected_error(self):
        case = self.browser_case("TC_FUNC_01")
        case.login_page.wait_for_error.side_effect = TimeoutException("No error response")
        with self.assertRaises(TimeoutException):
            case.test_TC_FUNC_01()

    def test_runner_spec_is_not_replaced_by_module_cached_excel_data(self):
        case = browser_tests.TestLogin("test_TC_SEC_02")
        case.settings = Settings(run_browser=True, run_security=True, username="test-user",
                                 invalid_password="wrong", error_selector="#error",
                                 error_text="Generic error", success_selector="#private")
        case.spec = copy.deepcopy(browser_tests.CASES["TC_SEC_02"])
        case.spec["parameters"]["payload"] = "RUNTIME_SQL_SAMPLE"
        page = Mock()
        page.wait_for_error.return_value.text = "Generic error"
        page.is_login_page.return_value = True
        page.get_page_source.return_value = "clean response"
        driver = Mock()
        driver.find_elements.return_value = []
        with patch.object(browser_tests.webdriver, "Chrome", return_value=driver), \
             patch.object(browser_tests, "LoginPage", return_value=page):
            result = unittest.TestResult()
            case.run(result)
        self.assertTrue(result.wasSuccessful(), str(result.errors) + str(result.failures))
        self.assertFalse(result.skipped)
        calls = page.login.call_args_list
        self.assertEqual(calls[0].args[0], "RUNTIME_SQL_SAMPLE")
        self.assertEqual(len(calls), 1)

    def test_sql_error_markers_do_not_match_benign_theme_or_product_names(self):
        case = self.browser_case("TC_SEC_06")
        case.assert_rejected = Mock()
        case.login_page.get_page_source.return_value = "<div class='flora-panel'>SQLServer documentation</div>"
        case.test_TC_SEC_06()
        case.login_page.get_page_source.return_value = "ORA-00933: SQL command not properly ended"
        with self.assertRaises(AssertionError):
            case.test_TC_SEC_06()

    def test_auth_rejection_does_not_submit_valid_password_or_access_protected_page(self):
        case = self.browser_case("TC_FUNC_01")
        case.settings = replace(case.settings, success_text="Test User", protected_url="https://example.test/private")
        case.login_page.wait_for_error.return_value.text = case.settings.error_text
        case.login_page.is_login_page.return_value = True
        case.test_TC_FUNC_01()
        case.driver.get.assert_not_called()
        case.login_page.wait_for_success.assert_not_called()
        case.login_page.login.assert_called_once_with(case.settings.username, case.settings.invalid_password, enter=False)

    def test_enter_cannot_pass_merely_because_login_button_remains(self):
        case = self.browser_case("TC_FUNC_10")
        case.login_page.wait_for_error.side_effect = TimeoutException("No submit")
        with self.assertRaises(TimeoutException):
            case.test_TC_FUNC_10()
        self.assertTrue(case.login_page.login.call_args.kwargs["enter"])

    def test_whitespace_case_cannot_pass_from_nonempty_input_alone(self):
        case = self.browser_case("TC_BND_01")
        case.settings = replace(case.settings, username_trim="accept")
        case.login_page.wait_for_error.side_effect = TimeoutException("No rejection")
        with self.assertRaises(TimeoutException):
            case.test_TC_BND_01()

    def test_checkbox_no_change_cannot_pass(self):
        case = self.browser_case("TC_FUNC_08")
        case.login_page.remember_checkbox.return_value.is_selected.return_value = False
        case.login_page.wait.until.side_effect = lambda condition: (
            True if condition(case.driver) else (_ for _ in ()).throw(TimeoutException("No change")))
        with self.assertRaises(TimeoutException):
            case.test_TC_FUNC_08()

    def test_missing_form_is_error_instead_of_vacuous_post_pass(self):
        settings = Settings()
        driver = Mock()
        driver.execute_script.return_value = None
        page = LoginPage(driver, settings)
        page.find = Mock()
        with self.assertRaises(ValueError):
            page.login_form()

    def test_absent_zero_negative_nonfinite_or_boolean_metric_fails(self):
        case = self.browser_case("TC_PERF_03")
        for metric in (None, {}, {"ttfb": 0}, {"ttfb": -1}, {"ttfb": float("nan")},
                       {"ttfb": float("inf")}, {"ttfb": True}):
            with self.subTest(metric=metric):
                case.driver.execute_script.return_value = metric
                with self.assertRaises(AssertionError):
                    case.navigation_metric("ttfb")

    def test_missing_error_or_missing_login_form_cannot_pass_negative_login(self):
        case = self.browser_case("TC_FUNC_05")
        case.login_page.wait_for_error.side_effect = TimeoutException("No error")
        with self.assertRaises(TimeoutException):
            case.test_TC_FUNC_05()
        case.login_page.wait_for_error.side_effect = None
        case.login_page.wait_for_error.return_value.text = "Generic error"
        case.settings = replace(case.settings, error_text="Generic error")
        case.login_page.is_login_page.return_value = True
        case.login_page.is_login_page.return_value = False
        with self.assertRaises(AssertionError):
            case.test_TC_FUNC_05()

    def test_xss_alert_and_unexpected_driver_errors_are_not_swallowed(self):
        case = self.browser_case("TC_SEC_07")
        case.login_page.wait.until.return_value = Mock()
        with self.assertRaises(AssertionError):
            case.test_TC_SEC_07()
        case.login_page.wait.until.side_effect = WebDriverException("driver disconnected")
        with self.assertRaises(WebDriverException):
            case.test_TC_SEC_07()

    def test_explicitly_disabled_browser_does_not_create_driver(self):
        case = browser_tests.TestLogin("test_TC_SEC_01")
        case.settings = Settings(run_browser=False)
        with patch.object(browser_tests.webdriver, "Chrome") as chrome:
            result = unittest.TestResult()
            case.run(result)
            chrome.assert_not_called()
            self.assertEqual(len(result.errors), 1)
            self.assertFalse(result.skipped)

    def test_missing_credentials_report_error_before_driver_creation(self):
        case = browser_tests.TestLogin("test_TC_FUNC_01")
        case.settings = Settings(run_browser=True, username="", password="")
        with patch.object(browser_tests.webdriver, "Chrome") as chrome:
            result = unittest.TestResult()
            case.run(result)
            chrome.assert_not_called()
            self.assertEqual(len(result.errors), 1)
            self.assertFalse(result.skipped)

    def test_defaults_start_visible_chrome_and_execute_browser_case(self):
        case = browser_tests.TestLogin("test_TC_SEC_01")
        page = Mock()
        page.get_attribute.side_effect = ["password", "MaskingTest_123"]
        driver = Mock()
        with patch.object(browser_tests.webdriver, "Chrome", return_value=driver) as chrome, \
             patch.object(browser_tests, "LoginPage", return_value=page):
            result = unittest.TestResult()
            case.run(result)
        self.assertTrue(result.wasSuccessful(), str(result.errors) + str(result.failures))
        self.assertFalse(result.skipped)
        options = chrome.call_args.kwargs["options"]
        self.assertFalse(any(argument.startswith("--headless") for argument in options.arguments))
        page.enter_password.assert_called_once_with("MaskingTest_123")
        driver.quit.assert_called_once()

    def test_driver_is_independent_and_quit_when_navigation_fails(self):
        drivers = [Mock(), Mock()]
        page = Mock()
        page.navigate.side_effect = WebDriverException("navigation failed")
        with patch.object(browser_tests.webdriver, "Chrome", side_effect=drivers) as chrome, \
             patch.object(browser_tests, "LoginPage", return_value=page):
            for _ in range(2):
                case = browser_tests.TestLogin("test_TC_SEC_01")
                case.settings = Settings(run_browser=True)
                result = unittest.TestResult()
                case.run(result)
                self.assertEqual(len(result.errors), 1)
            self.assertEqual(chrome.call_count, 2)
        for driver in drivers:
            driver.quit.assert_called_once()



    def test_report_keeps_skip_and_partial_success_distinct_and_redacts_secrets(self):
        def skip(case):
            case.skipTest("Prerequisite missing")
        def pass_check(case):
            case.assertTrue(True)
        def fail_check(case):
            case.fail("secret-pass must not leak")
        synthetic = type("Synthetic", (unittest.TestCase,), {
            "test_TC_FUNC_01": skip, "test_TC_FUNC_08": pass_check, "test_TC_UI_02": fail_check})
        cases = load_cases()
        next(case for case in cases if case['id'] == 'TC_FUNC_08')['automation'] = 'Partial'
        with tempfile.TemporaryDirectory() as directory:
            factory = lambda *a, **kw: EvidenceResult(*a, report_dir=Path(directory), cases=cases,
                                                     settings=Settings(password="secret-pass"), **kw)
            result = unittest.TextTestRunner(stream=io.StringIO(), resultclass=factory).run(
                unittest.defaultTestLoader.loadTestsFromTestCase(synthetic))
        self.assertEqual(result.records["TC_FUNC_01"]["status"], "SKIP")
        self.assertEqual(result.records["TC_FUNC_08"]["status"], "PARTIAL_PASS")
        self.assertEqual(result.records["TC_UI_02"]["status"], "FAIL")
        self.assertNotIn("secret-pass", json.dumps(result.records))
        self.assertNotIn("secret-pass", result.failures[0][1])
        self.assertEqual(result_exit_code(result), 1)

    def test_runner_exit_codes_do_not_treat_all_skipped_as_success(self):
        result = unittest.TestResult()
        self.assertEqual(result_exit_code(result), 2)
        result.testsRun = 1
        self.assertEqual(result_exit_code(result), 0)
        result.skipped.append((Mock(), "missing"))
        self.assertEqual(result_exit_code(result), 2)
        result.errors.append((Mock(), "error"))
        self.assertEqual(result_exit_code(result), 1)


if __name__ == "__main__":
    unittest.main(verbosity=2)
