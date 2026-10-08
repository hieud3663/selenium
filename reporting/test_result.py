"""Collect unittest outcomes and persist execution reports without altering testcase Excel."""
import json
import unittest
from math import ceil

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from config import SETTINGS
from testcases.reader import coverage_rows

class EvidenceResult(unittest.TextTestResult):
    def __init__(self, *args, report_dir, cases, allure_writer=None, settings=SETTINGS, **kwargs):
        super().__init__(*args, **kwargs)
        self.report_dir = report_dir
        self.cases = {c["id"]: c for c in cases}
        self.records = {}
        self.allure_writer = allure_writer
        self.settings = settings
        self.suite_errors = []

    def startTest(self, test):
        super().startTest(test)
        if self.allure_writer:
            id = test._testMethodName.removeprefix("test_")
            self.allure_writer.start(test, self.cases[id])

    def stopTest(self, test):
        if self.allure_writer:
            id = test._testMethodName.removeprefix("test_")
            record = self.records.get(id, {"status": "ERROR", "detail": "Test không có kết quả cuối."})
            self.allure_writer.finish(test, record, self.redact)
        super().stopTest(test)

    def redact(self, text):
        for name in ("username", "password", "invalid_password", "nonexistent_username"):
            secret = getattr(self.settings, name)
            if secret:
                text = text.replace(secret, "[REDACTED]")
        for case in self.cases.values():
            for name in ("password", "unicode_password"):
                secret = case["parameters"].get(name)
                if isinstance(secret, str) and secret:
                    text = text.replace(secret, "[REDACTED]")
        return text

    def _exc_info_to_string(self, err, test):
        return self.redact(super()._exc_info_to_string(err, test))

    def record(self, test, status, detail="", screenshot=False):
        if not hasattr(test, "_testMethodName"):
            self.suite_errors.append({"status": status, "detail": self.redact(detail)})
            return
        id = test._testMethodName.removeprefix("test_")
        existing = self.records.get(id)
        rank = {"PASS": 0, "PARTIAL_PASS": 0, "SKIP": 1, "FAIL": 2, "ERROR": 3}
        if existing and rank[existing["status"]] > rank[status]:
            return
        record = {"status": status, "detail": self.redact(detail), "screenshot": "",
                  "phase": getattr(test, "current_phase", "test")}
        if getattr(test, "observations", None):
            record["observations"] = test.observations
            record["detail"] += " " + json.dumps(test.observations)
        driver = getattr(test, "driver", None)
        if screenshot and driver:
            path = self.report_dir / f"{id}.png"
            try:
                driver.execute_script("document.querySelectorAll('input').forEach(e => e.value = '');")
                if driver.save_screenshot(str(path)):
                    record["screenshot"] = str(path)
            except Exception as error:
                record["evidence_error"] = self.redact(str(error))
        self.records[id] = record

    def addSuccess(self, test):
        super().addSuccess(test)
        id = test._testMethodName.removeprefix("test_")
        scope = self.cases[id]["automation"]
        self.record(test, "PARTIAL_PASS" if scope == "Partial" else "PASS",
                    "Các assertion đã triển khai đạt; xem Coverage để biết phạm vi còn lại.")

    def addFailure(self, test, err):
        super().addFailure(test, err)
        self.record(test, "FAIL", self._exc_info_to_string(err, test), screenshot=True)

    def addError(self, test, err):
        super().addError(test, err)
        summary = str(err[1]).split("Stacktrace:", 1)[0].strip().replace("\n", " ")
        self.stream.writeln(self.redact(f"  -> {getattr(test, 'current_phase', 'suite')}: {err[0].__name__}: {summary}"))
        self.record(test, "ERROR", self._exc_info_to_string(err, test), screenshot=True)

    def addSkip(self, test, reason):
        super().addSkip(test, reason)
        self.record(test, "SKIP", reason)

    def addSubTest(self, test, subtest, err):
        super().addSubTest(test, subtest, err)
        if err:
            status = "FAIL" if issubclass(err[0], test.failureException) else "ERROR"
            self.record(test, status, str(subtest) + "\n" + self._exc_info_to_string(err, test), screenshot=True)

def result_exit_code(result):
    if not result.wasSuccessful():
        return 1
    if result.testsRun == 0 or result.skipped:
        return 2
    return 0


def save_json(report_dir, report):
    (report_dir / "results.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def format_sheet(sheet, widths):
    sheet.freeze_panes = "B2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1A365D")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    for row in sheet.iter_rows(min_row=2):
        lines = 1
        for cell, width in zip(row, widths):
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            paragraphs = str(cell.value or "").splitlines() or [""]
            lines = max(lines, sum(max(1, ceil(len(paragraph) / width)) for paragraph in paragraphs))
        sheet.row_dimensions[row[0].row].height = min(409, lines * 16 + 10)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def save_results(report_dir, cases, report):
    save_json(report_dir, report)
    workbook = Workbook()
    execution = workbook.active
    execution.title = "Execution"
    execution.append(["Test Case ID", "Kết quả chạy", "Phạm vi", "Quan sát", "Bằng chứng", "Run ID"])
    for case in cases:
        record = report["results"].get(case["id"], {})
        detail = record.get("detail", "Chưa chạy; xem Coverage.")
        if len(detail) > 800:
            detail = detail[:800] + "\nChi tiết đầy đủ trong results.json và Allure."
        execution.append([case["id"], record.get("status", "Not Run"), case["automation"],
                          detail, record.get("screenshot", ""), report["run_id"]])
    format_sheet(execution, [17, 20, 18, 90, 60, 36])
    coverage = workbook.create_sheet("Coverage")
    coverage.append(["Test Case ID", "Phạm vi automation", "Phương thức", "Kiểm tra đã triển khai", "Kiểm tra còn lại"])
    for row in coverage_rows(cases):
        coverage.append(row)
    format_sheet(coverage, [17, 20, 54, 66, 72])
    workbook.save(report_dir / "results.xlsx")
