"""Read testcase specifications directly from Excel without changing the workbook."""
import ast
import json
from pathlib import Path

from openpyxl import load_workbook
from config import ROOT

WORKBOOK_PATH = ROOT / "testcases" / "TestCases_VanPhongDienTu_UTC.xlsx"
FIELDS = {
    "id": "Test Case ID", "module": "Phân hệ (Module)",
    "type": "Loại kiểm thử (Type)", "title": "Mô tả kịch bản kiểm thử",
    "preconditions": "Tiền điều kiện", "steps": "Các bước thực hiện",
    "test_data": "Dữ liệu kiểm thử", "expected": "Kết quả mong đợi",
    "postconditions": "Hậu điều kiện", "actual": "Kết quả thực tế",
    "status": "Trạng thái", "severity": "Mức độ ảnh hưởng nếu lỗi",
    "priority": "Ưu tiên", "requirement": "Căn cứ và yêu cầu cần xác minh",
    "automation": "Phạm vi automation", "automated_checks": "Kiểm tra đã triển khai",
    "manual_checks": "Kiểm tra còn lại", "parameters": "Tham số automation (JSON)",
}


def validate_cases(cases):
    expected_ids = {f"{prefix}_{i:02}" for prefix, count in
                    [("TC_FUNC", 12), ("TC_BND", 5), ("TC_SEC", 12), ("TC_PERF", 6), ("TC_UI", 5)]
                    for i in range(1, count + 1)}
    ids = [case.get("id") for case in cases]
    if len(ids) != 40 or set(ids) != expected_ids:
        raise ValueError("Excel phải có đúng 40 ID gốc, không thiếu hoặc trùng.")
    for case in cases:
        case_id = case["id"]
        if set(case) != set(FIELDS):
            raise ValueError(f"{case_id}: thiếu hoặc thừa thuộc tính testcase.")
        if not all(isinstance(case[key], str) for key in FIELDS if key != "parameters"):
            raise ValueError(f"{case_id}: các cột mô tả phải là chuỗi.")
        if case["status"] != "Not Run" or case["actual"]:
            raise ValueError("Excel là đặc tả: Status=Not Run, Actual trống; kết quả chạy lưu trong reports.")
        if case["automation"] not in {"Automated", "Partial", "Manual"}:
            raise ValueError(f"{case_id}: phạm vi automation không hợp lệ.")
        if not all(case[key].strip() for key in ("title", "preconditions", "steps", "expected", "requirement")):
            raise ValueError(f"{case_id}: đặc tả bắt buộc không được trống.")
        parameters = case["parameters"]
        if not isinstance(parameters, dict):
            raise ValueError(f"{case_id}: tham số phải là JSON object.")
        if "auth_flow" in parameters and not isinstance(parameters["auth_flow"], bool):
            raise ValueError(f"{case_id}: auth_flow phải là boolean.")
        if "payload" in parameters and (not isinstance(parameters["payload"], str) or not parameters["payload"]):
            raise ValueError(f"{case_id}: payload phải là chuỗi không rỗng.")
        if case_id in {"TC_FUNC_02", "TC_FUNC_03", "TC_FUNC_04"}:
            field = parameters.get("empty_field")
            if not isinstance(field, str) or field not in {"username", "password", "both"}:
                raise ValueError(f"{case_id}: empty_field không hợp lệ.")
            if not isinstance(parameters.get("error_text"), str) or not parameters["error_text"].strip():
                raise ValueError(f"{case_id}: thiếu error_text cho phản hồi trường rỗng.")
        if case_id in {"TC_SEC_02", "TC_SEC_03", "TC_SEC_04", "TC_SEC_06", "TC_SEC_07", "TC_SEC_08"}:
            if not parameters.get("payload"):
                raise ValueError(f"{case_id}: thiếu payload.")
        if case_id == "TC_UI_03":
            order = parameters.get("tab_order")
            if not isinstance(order, list) or len(order) < 2 or any(
                    not isinstance(name, str) or name not in {"username", "password", "login", "remember", "email"}
                    for name in order):
                raise ValueError("TC_UI_03: tab_order không hợp lệ.")
    return cases


def auth_flow_ids(cases):
    """The full 40-case auth suite; no hidden subset or security opt-in."""
    return [case["id"] for case in cases]


def load_cases(path=None):
    source = Path(path) if path is not None else WORKBOOK_PATH
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        if "Testcases" not in workbook.sheetnames:
            raise ValueError("Excel thiếu sheet Testcases.")
        rows = workbook["Testcases"].iter_rows()
        header = next(rows, ())
        if [cell.value for cell in header] != list(FIELDS.values()):
            raise ValueError("Header Excel không khớp 18 cột của template.")
        cases = []
        for row_number, row in enumerate(rows, 2):
            if all(cell.value is None for cell in row):
                continue
            if any(cell.data_type == "f" for cell in row):
                raise ValueError(f"Dòng {row_number}: testcase phải là dữ liệu tĩnh, không dùng công thức.")
            values = ["" if cell.value is None else cell.value for cell in row]
            case = dict(zip(FIELDS, values))
            try:
                case["parameters"] = json.loads(case["parameters"])
            except (json.JSONDecodeError, TypeError) as error:
                raise ValueError(f"Dòng {row_number} ({case.get('id')}): JSON tham số không hợp lệ.") from error
            cases.append(case)
        return validate_cases(cases)
    finally:
        workbook.close()


def validate_mapping(cases, test_path=None):
    source = Path(test_path) if test_path is not None else ROOT / "tests" / "test_login.py"
    tree = ast.parse(source.read_text(encoding="utf-8"))
    cls = next((node for node in tree.body if isinstance(node, ast.ClassDef) and node.name == "TestLogin"), None)
    if cls is None:
        raise ValueError("Không tìm thấy lớp TestLogin để đối chiếu ma trận.")
    methods = {node.name for node in cls.body if isinstance(node, ast.FunctionDef) and node.name.startswith("test_")}
    expected = {"test_" + case["id"] for case in cases if case["automation"] != "Manual"}
    if methods != expected:
        raise ValueError(f"Excel không khớp code: thiếu {sorted(expected - methods)}, thừa {sorted(methods - expected)}")


def coverage_rows(cases):
    return [[case["id"], case["automation"],
             f"tests.test_login.TestLogin.test_{case['id']}" if case["automation"] != "Manual" else "",
             case["automated_checks"], case["manual_checks"]] for case in cases]
